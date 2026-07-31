#!/usr/bin/env python3
"""Generate Crownsmen campaign performance report PPTX from collected JSON."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


FOOTER = "© 2018 Crownsmen Partners | www.crownsmen.com | info@crownsmen.com"


def fmt_num(value) -> str:
    if value is None:
        return "[TBD]"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def add_title_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    report = data.get("report") or {}
    year = report.get("year", "")
    date_created = report.get("date_created", "")

    title = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1.5))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = f"Campaign Performance Report {year}"
    p.font.size = Pt(36)
    p.font.bold = True

    footer = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(0.5))
    footer.text_frame.paragraphs[0].text = FOOTER
    footer.text_frame.paragraphs[0].font.size = Pt(10)

    date_box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.4))
    date_box.text_frame.paragraphs[0].text = f"Date Created {date_created}"
    date_box.text_frame.paragraphs[0].font.size = Pt(12)


def add_text_slide(prs: Presentation, heading: str, body: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.5), Inches(0.8))
    title.text_frame.paragraphs[0].text = heading
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.bold = True

    if body:
        content = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.5), Inches(5))
        content.text_frame.paragraphs[0].text = body
        content.text_frame.paragraphs[0].font.size = Pt(16)


def add_table_slide(prs: Presentation, heading: str, headers: list[str], rows: list[list[str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.5), Inches(0.7))
    title.text_frame.paragraphs[0].text = heading
    title.text_frame.paragraphs[0].font.size = Pt(24)
    title.text_frame.paragraphs[0].font.bold = True

    cols = len(headers)
    table_shape = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.5), Inches(1.4), Inches(9), Inches(0.4 * (len(rows) + 2)))
    table = table_shape.table

    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(11)

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value or ""


def add_youtube_stats_slide(prs: Presentation, heading: str, yt: dict | None) -> None:
    yt = yt or {}
    subtitle = yt.get("title") or ""
    metrics = [
        ("Video Views", fmt_num(yt.get("view_count"))),
        ("Average Percentage Viewed", yt.get("average_percentage_viewed") or "[TBD]"),
        ("Average View Duration", yt.get("average_view_duration") or "[TBD]"),
    ]
    if "Full Episode" in heading:
        metrics.append(("Watch Time Hours", fmt_num(yt.get("watch_time_hours"))))
    metrics.extend(
        [
            ("% of viewers watching 25% or more", yt.get("pct_25_plus") or "[TBD]"),
            ("% of viewers watching 50% or more", yt.get("pct_50_plus") or "[TBD]"),
            ("% of viewers watching 90% or more", yt.get("pct_90_plus") or "[TBD]"),
        ]
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.6))
    title.text_frame.paragraphs[0].text = heading
    title.text_frame.paragraphs[0].font.size = Pt(22)
    title.text_frame.paragraphs[0].font.bold = True

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(8.5), Inches(0.8))
    sub.text_frame.paragraphs[0].text = subtitle
    sub.text_frame.paragraphs[0].font.size = Pt(14)

    y = 2.0
    for label, value in metrics:
        lbl = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(5.5), Inches(0.35))
        lbl.text_frame.paragraphs[0].text = label
        lbl.text_frame.paragraphs[0].font.size = Pt(14)
        val = slide.shapes.add_textbox(Inches(6.5), Inches(y), Inches(2.5), Inches(0.35))
        val.text_frame.paragraphs[0].text = value
        val.text_frame.paragraphs[0].font.size = Pt(14)
        val.text_frame.paragraphs[0].font.bold = True
        y += 0.45


def add_geography_slide(prs: Presentation, yt: dict | None) -> None:
    yt = yt or {}
    geo = yt.get("geography") or []
    rows = [[g.get("country", ""), fmt_num(g.get("views"))] for g in geo[:10]]
    if not rows:
        rows = [["[TBD]", "Pull from YouTube Analytics API or Studio export"]]
    add_table_slide(prs, "YouTube | Geography", ["Country/Region", "Views"], rows)


def add_email_conversion_slide(prs: Presentation, data: dict) -> None:
    email = data.get("email_marketing") or {}
    conv = data.get("one_click_conversion") or {}

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    t1 = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(4), Inches(0.5))
    t1.text_frame.paragraphs[0].text = "Email Marketing"
    t1.text_frame.paragraphs[0].font.size = Pt(20)
    t1.text_frame.paragraphs[0].font.bold = True

    email_metrics = [
        ("Successful Deliveries", fmt_num(email.get("successful_deliveries"))),
        ("Opens", fmt_num(email.get("opens"))),
        ("Open Rate", email.get("open_rate") or "[TBD]"),
        ("Clicks", fmt_num(email.get("clicks"))),
    ]
    y = 1.2
    for label, value in email_metrics:
        slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(3.5), Inches(0.35)).text_frame.paragraphs[0].text = label
        slide.shapes.add_textbox(Inches(4.5), Inches(y), Inches(1.5), Inches(0.35)).text_frame.paragraphs[0].text = value
        y += 0.4

    t2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8), Inches(0.5))
    t2.text_frame.paragraphs[0].text = "One-Click Conversion Results"
    t2.text_frame.paragraphs[0].font.size = Pt(20)
    t2.text_frame.paragraphs[0].font.bold = True

    slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(2), Inches(0.35)).text_frame.paragraphs[0].text = "URL Used"
    url_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.5), Inches(0.6))
    url_box.text_frame.paragraphs[0].text = conv.get("tracking_url") or "[TBD]"
    url_box.text_frame.paragraphs[0].font.size = Pt(10)

    slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(3), Inches(0.35)).text_frame.paragraphs[0].text = "Total # Clicks to Website"
    slide.shapes.add_textbox(Inches(4.5), Inches(5.0), Inches(2), Inches(0.35)).text_frame.paragraphs[0].text = fmt_num(conv.get("total_clicks"))


def generate(data: dict, output_path: Path) -> None:
    prs = Presentation()
    report = data.get("report") or {}
    ep_num = report.get("episode_number", "01")

    add_title_slide(prs, data)
    add_text_slide(prs, "Report Content", "Episode Performance 01")

    launch = data.get("episode_launch_links") or {}
    launch_rows = [[platform, url or ""] for platform, url in launch.items()]
    add_table_slide(prs, "Promotional Links - Episode Launch", ["Platform Name", "URL"], launch_rows)

    event_rows = [
        [item.get("content_type", ""), item.get("link", ""), item.get("remarks", "")]
        for item in (data.get("event_promotion") or [])
    ]
    add_table_slide(prs, "Promotional Links - Event", ["Content Type", "Link", "Remarks"], event_rows)

    add_text_slide(prs, "Episode Performance", str(ep_num))
    add_youtube_stats_slide(prs, "YouTube | Full Episode Statistics", data.get("youtube_full"))
    add_youtube_stats_slide(prs, "YouTube | Highlight Episode Statistics", data.get("youtube_highlight"))
    add_geography_slide(prs, data.get("youtube_full"))

    add_email_conversion_slide(prs, data)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    thanks = slide.shapes.add_textbox(Inches(2.5), Inches(3), Inches(5), Inches(1))
    thanks.text_frame.paragraphs[0].text = "Thank You!"
    thanks.text_frame.paragraphs[0].font.size = Pt(36)
    thanks.text_frame.paragraphs[0].font.bold = True
    footer = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(0.5))
    footer.text_frame.paragraphs[0].text = FOOTER
    footer.text_frame.paragraphs[0].font.size = Pt(10)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main() -> None:
    json_path = Path(sys.argv[1] if len(sys.argv) > 1 else "data/collected.json")
    out_path = Path(sys.argv[2] if len(sys.argv) > 2 else "output/report.pptx")

    with json_path.open(encoding="utf-8") as fh:
        data = json.load(fh)

    generate(data, out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
