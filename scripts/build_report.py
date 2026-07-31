#!/usr/bin/env python3
"""Build Crownsmen Partners Campaign Performance Report (PPTX + PDF)."""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Add scripts dir to path for imports
sys.path.insert(0, str(Path(__file__).resolve().parent))
from fetch_metrics import CampaignMetrics, fetch_all, fmt_num, metrics_to_dict

# Crownsmen brand colors
NAVY = RGBColor(0x1A, 0x2B, 0x4A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT_BLUE = RGBColor(0x2E, 0x6B, 0xA4)


def val(v: int | None, suffix: str = "") -> str:
    if v is None:
        return "[TBD]"
    return f"{v:,}{suffix}"


def pct(v: float | None) -> str:
    if v is None:
        return "[TBD]"
    return f"{v:.1f}%"


def money(v: float | None) -> str:
    if v is None:
        return "[TBD]"
    return f"${v:,.2f}"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title: str, subtitle: str = "") -> None:
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.1))  # MSO_SHAPE.RECTANGLE
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.6))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(0.65), Inches(9), Inches(0.35))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = GOLD


def add_footer(slide, text: str) -> None:
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(9), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = MID_GRAY
    p.alignment = PP_ALIGN.RIGHT


def add_metric_box(slide, left, top, width, height, label: str, value: str, highlight: bool = False) -> None:
    box = slide.shapes.add_shape(1, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = GOLD if highlight else LIGHT_GRAY
    box.line.color.rgb = GOLD if highlight else RGBColor(0xDD, 0xDD, 0xDD)

    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), height - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(11)
    p1.font.color.rgb = MID_GRAY if not highlight else NAVY
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(22 if highlight else 18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY


def add_table(slide, left, top, width, rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    n_rows = len(rows)
    n_cols = len(rows[0]) if rows else 0
    if n_rows == 0:
        return

    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.35 * n_rows))
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)

    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = cell_text
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(10)
                paragraph.font.color.rgb = WHITE if r == 0 else DARK_GRAY
                if r == 0:
                    paragraph.font.bold = True
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY


def build_presentation(metrics: CampaignMetrics) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    campaign = metrics.campaign
    client = campaign.get("client", "Client")
    episode = campaign.get("episode", "")
    title = campaign.get("title", "")
    series = campaign.get("series", "")
    report_date = campaign.get("report_date", datetime.now().strftime("%Y-%m-%d"))
    footer = f"Crownsmen Partners | Confidential | {report_date}"

    # --- Slide 1: Cover ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.5), Inches(8.5), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Campaign Performance Report"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE

    tb2 = slide.shapes.add_textbox(Inches(0.75), Inches(2.8), Inches(8.5), Inches(1.5))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = f"{episode}. {client}"
    p2.font.size = Pt(28)
    p2.font.color.rgb = GOLD
    p3 = tf2.add_paragraph()
    p3.text = title
    p3.font.size = Pt(20)
    p3.font.color.rgb = WHITE

    tb3 = slide.shapes.add_textbox(Inches(0.75), Inches(4.8), Inches(8.5), Inches(1.0))
    tf3 = tb3.text_frame
    p4 = tf3.paragraphs[0]
    p4.text = series
    p4.font.size = Pt(16)
    p4.font.color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    p5 = tf3.add_paragraph()
    p5.text = f"Report Date: {report_date}"
    p5.font.size = Pt(14)
    p5.font.color.rgb = RGBColor(0xAA, 0xBB, 0xCC)

    # --- Slide 2: Executive Summary ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Executive Summary", f"{episode} — {client}")

    main_yt = next((y for y in metrics.youtube if y.extra.get("type") == "main"), metrics.youtube[0] if metrics.youtube else None)
    total_yt_views = sum(y.views or 0 for y in metrics.youtube)
    total_li_reactions = sum(s.reactions or 0 for s in metrics.social if s.platform == "LinkedIn")
    email = metrics.email
    ads = metrics.google_ads
    analytics = metrics.analytics

    add_metric_box(slide, Inches(0.5), Inches(1.4), Inches(2.1), Inches(1.1), "Total YouTube Views", val(total_yt_views), True)
    add_metric_box(slide, Inches(2.8), Inches(1.4), Inches(2.1), Inches(1.1), "Main Episode Views", val(main_yt.views if main_yt else None))
    add_metric_box(slide, Inches(5.1), Inches(1.4), Inches(2.1), Inches(1.1), "LinkedIn Reactions", val(total_li_reactions))
    add_metric_box(slide, Inches(7.4), Inches(1.4), Inches(2.1), Inches(1.1), "Email Open Rate", pct(email.get("open_rate")))

    add_metric_box(slide, Inches(0.5), Inches(2.7), Inches(2.1), Inches(1.1), "Email Sends", val(email.get("sends")))
    add_metric_box(slide, Inches(2.8), Inches(2.7), Inches(2.1), Inches(1.1), "Email Clicks", val(email.get("clicks")))
    add_metric_box(slide, Inches(5.1), Inches(2.7), Inches(2.1), Inches(1.1), "Landing Page Sessions", val(analytics.get("sessions")))
    add_metric_box(slide, Inches(7.4), Inches(2.7), Inches(2.1), Inches(1.1), "UTM Clicks", val(analytics.get("utm_clicks")))

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.1), Inches(9), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Campaign Overview"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    overview_lines = [
        f"• Full episode published across YouTube, podcast platforms, and social channels",
        f"• {len([y for y in metrics.youtube if y.extra.get('type') == 'clip'])} promotional clips distributed on YouTube",
        f"• Social promotion on LinkedIn, Facebook, Instagram, Threads, and X",
        f"• Email campaign via Constant Contact to industry subscriber list",
        f"• Dedicated landing page: crownsmen.com/equify-financial-construction-equipment-financing/",
    ]
    for line in overview_lines:
        bp = tf.add_paragraph()
        bp.text = line
        bp.font.size = Pt(12)
        bp.font.color.rgb = DARK_GRAY
        bp.space_before = Pt(4)

    add_footer(slide, footer)

    # --- Slide 3: YouTube Performance ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "YouTube Performance", "Crownsmen Partners Channel")

    yt_rows = [["Video", "Type", "Views", "Likes", "Comments", "Duration"]]
    for yt in metrics.youtube:
        dur = "[TBD]"
        if yt.duration_sec:
            m, s = divmod(yt.duration_sec, 60)
            h, m = divmod(m, 60)
            dur = f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"
        short_title = (yt.title or "Untitled")[:45] + ("..." if yt.title and len(yt.title) > 45 else "")
        yt_rows.append([
            short_title,
            yt.extra.get("type", "").title(),
            val(yt.views),
            val(yt.likes),
            val(yt.comments),
            dur,
        ])

    add_table(slide, Inches(0.4), Inches(1.3), Inches(9.2), yt_rows, [3.5, 0.8, 1.0, 0.8, 0.9, 0.8])

    if main_yt and main_yt.views:
        clip_views = sum(y.views or 0 for y in metrics.youtube if y.extra.get("type") == "clip")
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(2.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = "YouTube Highlights"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY
        highlights = [
            f"Main episode: {val(main_yt.views)} views",
            f"Promotional clips combined: {val(clip_views)} views",
            f"Total YouTube reach: {val(total_yt_views)} views",
        ]
        for h in highlights:
            hp = tf.add_paragraph()
            hp.text = f"• {h}"
            hp.font.size = Pt(12)
            hp.font.color.rgb = DARK_GRAY

    add_footer(slide, footer)

    # --- Slide 4: Social Media — LinkedIn ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Social Media — LinkedIn", "Organic Posts")

    li_rows = [["Account", "Reactions", "Comments", "Impressions", "Clicks", "Reposts"]]
    for s in metrics.social:
        if s.platform == "LinkedIn":
            li_rows.append([
                s.account or "LinkedIn",
                val(s.reactions),
                val(s.comments),
                val(s.impressions),
                val(s.clicks),
                val(s.shares),
            ])
    add_table(slide, Inches(0.4), Inches(1.5), Inches(9.2), li_rows, [2.5, 1.2, 1.2, 1.5, 1.2, 1.2])

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(9), Inches(2.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Post Details"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    for s in metrics.social:
        if s.platform == "LinkedIn":
            lp = tf.add_paragraph()
            lp.text = f"• {s.account}: {s.url}"
            lp.font.size = Pt(10)
            lp.font.color.rgb = ACCENT_BLUE

    add_footer(slide, footer)

    # --- Slide 5: Social Media — Facebook, Instagram, Threads, X ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Social Media — Other Platforms", "Facebook | Instagram | Threads | X")

    soc_rows = [["Platform", "Account", "Likes/Reactions", "Comments", "Shares", "Views/Reach"]]
    for s in metrics.social:
        if s.platform != "LinkedIn":
            likes = s.likes or s.reactions
            reach = s.views or s.impressions
            soc_rows.append([
                s.platform,
                s.account or "—",
                val(likes),
                val(s.comments),
                val(s.shares),
                val(reach),
            ])
    add_table(slide, Inches(0.4), Inches(1.5), Inches(9.2), soc_rows, [1.5, 2.0, 1.5, 1.2, 1.2, 1.5])

    add_footer(slide, footer)

    # --- Slide 6: Podcast & Audio Distribution ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Podcast & Audio Distribution", series)

    pod_rows = [["Platform", "Episode Title", "Plays/Downloads", "Duration"]]
    for p in metrics.podcast:
        dur = "[TBD]"
        if p.duration_sec:
            dur = f"{p.duration_sec // 60}m"
        elif p.extra.get("description"):
            m = __import__("re").search(r"(\d+)m", p.extra["description"])
            if m:
                dur = f"{m.group(1)}m"
        short = (p.title or "CS 67 — Equify Financial")[:40]
        pod_rows.append([
            p.platform,
            short,
            val(p.plays or p.views),
            dur,
        ])
    add_table(slide, Inches(0.4), Inches(1.5), Inches(9.2), pod_rows, [1.8, 3.5, 1.5, 1.0])

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(2.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Distribution Notes"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    notes = [
        "Episode syndicated to Spotify, Apple Podcasts, Amazon Music, and Rumble",
        "Full video also available on YouTube and Crownsmen.com landing page",
    ]
    for n in notes:
        np = tf.add_paragraph()
        np.text = f"• {n}"
        np.font.size = Pt(12)
        np.font.color.rgb = DARK_GRAY

    add_footer(slide, footer)

    # --- Slide 7: Email Marketing (Constant Contact) ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Email Marketing", "Constant Contact")

    add_metric_box(slide, Inches(0.5), Inches(1.5), Inches(2.1), Inches(1.2), "Emails Sent", val(email.get("sends")), True)
    add_metric_box(slide, Inches(2.8), Inches(1.5), Inches(2.1), Inches(1.2), "Delivered", val(email.get("delivered")))
    add_metric_box(slide, Inches(5.1), Inches(1.5), Inches(2.1), Inches(1.2), "Unique Opens", val(email.get("unique_opens")))
    add_metric_box(slide, Inches(7.4), Inches(1.5), Inches(2.1), Inches(1.2), "Open Rate", pct(email.get("open_rate")))

    add_metric_box(slide, Inches(0.5), Inches(3.0), Inches(2.1), Inches(1.2), "Unique Clicks", val(email.get("unique_clicks")))
    add_metric_box(slide, Inches(2.8), Inches(3.0), Inches(2.1), Inches(1.2), "Click Rate", pct(email.get("click_rate")))
    add_metric_box(slide, Inches(5.1), Inches(3.0), Inches(2.1), Inches(1.2), "Bounces", val(email.get("bounces")))
    add_metric_box(slide, Inches(7.4), Inches(3.0), Inches(2.1), Inches(1.2), "Unsubscribes", val(email.get("unsubscribes")))

    email_rows = [["Metric", "Value"]]
    for k, label in [
        ("campaign_name", "Campaign Name"),
        ("subject_line", "Subject Line"),
        ("send_date", "Send Date"),
    ]:
        email_rows.append([label, email.get(k) or "[TBD]"])
    add_table(slide, Inches(0.5), Inches(4.5), Inches(5), email_rows, [2.0, 3.0])

    add_footer(slide, footer)

    # --- Slide 8: Website & UTM Analytics ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Website & Landing Page Analytics", "crownsmen.com | UTM Tracking")

    add_metric_box(slide, Inches(0.5), Inches(1.5), Inches(2.1), Inches(1.2), "Page Views", val(analytics.get("pageviews")), True)
    add_metric_box(slide, Inches(2.8), Inches(1.5), Inches(2.1), Inches(1.2), "Sessions", val(analytics.get("sessions")))
    add_metric_box(slide, Inches(5.1), Inches(1.5), Inches(2.1), Inches(1.2), "Unique Users", val(analytics.get("users")))
    add_metric_box(slide, Inches(7.4), Inches(1.5), Inches(2.1), Inches(1.2), "Avg. Duration", analytics.get("avg_duration") or "[TBD]")

    add_metric_box(slide, Inches(0.5), Inches(3.0), Inches(2.1), Inches(1.2), "UTM Clicks", val(analytics.get("utm_clicks")))
    add_metric_box(slide, Inches(2.8), Inches(3.0), Inches(2.1), Inches(1.2), "Bounce Rate", pct(analytics.get("bounce_rate")))
    add_metric_box(slide, Inches(5.1), Inches(3.0), Inches(2.1), Inches(1.2), "Client Site Clicks", val(analytics.get("client_clicks")))
    add_metric_box(slide, Inches(7.4), Inches(3.0), Inches(2.1), Inches(1.2), "Conversion Rate", pct(analytics.get("conversion_rate")))

    utm_rows = [["UTM Source", "UTM Medium", "UTM Campaign", "Clicks"]]
    utm_data = analytics.get("utm_breakdown", [])
    if utm_data:
        for row in utm_data:
            utm_rows.append([row.get("source", "—"), row.get("medium", "—"), row.get("campaign", "—"), val(row.get("clicks"))])
    else:
        utm_rows.append(["[TBD]", "[TBD]", campaign.get("utm_campaign", "crownsmen"), "[TBD]"])
    add_table(slide, Inches(0.5), Inches(4.5), Inches(9), utm_rows, [2.0, 2.0, 2.5, 1.5])

    add_footer(slide, footer)

    # --- Slide 9: Google Ads Performance ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Google Ads Performance", "Paid Media")

    add_metric_box(slide, Inches(0.5), Inches(1.5), Inches(2.1), Inches(1.2), "Impressions", val(ads.get("impressions")), True)
    add_metric_box(slide, Inches(2.8), Inches(1.5), Inches(2.1), Inches(1.2), "Clicks", val(ads.get("clicks")))
    add_metric_box(slide, Inches(5.1), Inches(1.5), Inches(2.1), Inches(1.2), "CTR", pct(ads.get("ctr")))
    add_metric_box(slide, Inches(7.4), Inches(1.5), Inches(2.1), Inches(1.2), "Cost", money(ads.get("cost")))

    add_metric_box(slide, Inches(0.5), Inches(3.0), Inches(2.1), Inches(1.2), "Conversions", val(ads.get("conversions")))
    add_metric_box(slide, Inches(2.8), Inches(3.0), Inches(2.1), Inches(1.2), "Cost/Conv.", money(ads.get("cost_per_conversion")))
    add_metric_box(slide, Inches(5.1), Inches(3.0), Inches(2.1), Inches(1.2), "Avg. CPC", money(ads.get("avg_cpc")))
    add_metric_box(slide, Inches(7.4), Inches(3.0), Inches(2.1), Inches(1.2), "Video Views", val(ads.get("video_views")))

    ads_rows = [["Campaign", "Ad Group", "Impressions", "Clicks", "Cost"]]
    ads_breakdown = ads.get("campaigns", [])
    if ads_breakdown:
        for row in ads_breakdown:
            ads_rows.append([
                row.get("name", "—"),
                row.get("ad_group", "—"),
                val(row.get("impressions")),
                val(row.get("clicks")),
                money(row.get("cost")),
            ])
    else:
        ads_rows.append(["[TBD]", "[TBD]", "[TBD]", "[TBD]", "[TBD]"])
    add_table(slide, Inches(0.5), Inches(4.5), Inches(9), ads_rows, [2.5, 2.0, 1.5, 1.2, 1.5])

    add_footer(slide, footer)

    # --- Slide 10: Summary & Key Takeaways ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "Summary & Key Takeaways", f"{client} — {episode}")

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(5.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Performance Highlights"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    highlights = []
    if main_yt and main_yt.views:
        highlights.append(f"Main YouTube episode reached {main_yt.views:,} views")
    if total_yt_views:
        highlights.append(f"Combined YouTube content: {total_yt_views:,} total views")
    if total_li_reactions:
        highlights.append(f"LinkedIn posts generated {total_li_reactions} reactions")
    clip_count = len([y for y in metrics.youtube if y.extra.get("type") == "clip"])
    if clip_count:
        highlights.append(f"{clip_count} promotional clips extending episode reach")
    highlights.append("Multi-platform distribution across 10+ channels")
    if not highlights:
        highlights.append("Campaign distributed across YouTube, social, podcast, and email channels")

    for h in highlights:
        hp = tf.add_paragraph()
        hp.text = f"✓  {h}"
        hp.font.size = Pt(13)
        hp.font.color.rgb = DARK_GRAY
        hp.space_before = Pt(6)

    p2_title = tf.add_paragraph()
    p2_title.text = "\nData Gaps"
    p2_title.font.size = Pt(14)
    p2_title.font.bold = True
    p2_title.font.color.rgb = NAVY

    gaps = metrics.missing[:5] if metrics.missing else ["All key metrics collected"]
    for gap in gaps:
        gp = tf.add_paragraph()
        gp.text = f"• {gap[:80]}{'...' if len(gap) > 80 else ''}"
        gp.font.size = Pt(10)
        gp.font.color.rgb = MID_GRAY
        gp.space_before = Pt(3)

    # Summary metrics panel
    add_metric_box(slide, Inches(6.3), Inches(1.6), Inches(3.2), Inches(1.0), "Total YouTube Views", val(total_yt_views), True)
    add_metric_box(slide, Inches(6.3), Inches(2.8), Inches(3.2), Inches(1.0), "LinkedIn Reactions", val(total_li_reactions))
    add_metric_box(slide, Inches(6.3), Inches(4.0), Inches(3.2), Inches(1.0), "Email Open Rate", pct(email.get("open_rate")))
    add_metric_box(slide, Inches(6.3), Inches(5.2), Inches(3.2), Inches(1.0), "Landing Page Sessions", val(analytics.get("sessions")))

    add_footer(slide, footer)

    return prs


def write_missing_data(metrics: CampaignMetrics, output_path: Path) -> None:
    lines = [
        "# Missing Data Report",
        "",
        f"Generated: {metrics.fetched_at}",
        f"Campaign: {metrics.campaign.get('episode', '')} — {metrics.campaign.get('client', '')}",
        "",
        "The following metrics could not be automatically retrieved and are marked as **[TBD]** in the report deck.",
        "",
    ]

    if not metrics.missing:
        lines.append("All requested metrics were successfully retrieved.")
    else:
        lines.append("## Missing Metrics")
        lines.append("")
        for i, item in enumerate(metrics.missing, 1):
            lines.append(f"{i}. {item}")
        lines.append("")

    lines.extend([
        "## How to Fill Gaps",
        "",
        "### Constant Contact (Email)",
        "- Authenticate Zapier MCP in Cursor desktop (Settings → Tools & MCP → Connect)",
        "- Enable Constant Contact read actions in Zapier MCP",
        "- Re-run: `python3 scripts/build_report.py links.yaml -o output`",
        "",
        "### Google Ads",
        "- Set `GOOGLE_ADS_DEVELOPER_TOKEN`, `GOOGLE_ADS_CLIENT_ID`, `GOOGLE_ADS_CLIENT_SECRET`, `GOOGLE_ADS_REFRESH_TOKEN`",
        "- Or export metrics to `output/google_ads.json` and set `GOOGLE_ADS_METRICS_FILE`",
        "",
        "### Google Analytics (UTM Clicks)",
        "- Set `GA4_PROPERTY_ID` and `GOOGLE_APPLICATION_CREDENTIALS`",
        "- Or export UTM click data to `output/analytics.json` and set `GA_METRICS_FILE`",
        "",
        "### YouTube Analytics API",
        "- Set `YOUTUBE_API_KEY` for public stats (views, likes, comments)",
        "- YouTube Analytics API (watch time, traffic sources) requires OAuth credentials",
        "",
        "### Social Platforms",
        "- LinkedIn, Facebook, Instagram, Threads, X require platform API credentials or creator dashboard exports",
        "",
        "### Podcast Platforms",
        "- Spotify for Podcasters, Apple Podcasts Connect, Amazon Music, Rumble creator dashboards",
        "",
        "## Successfully Retrieved",
        "",
    ])

    lines.append("### YouTube")
    for yt in metrics.youtube:
        status = "✓" if yt.views else "⚠"
        lines.append(f"- {status} {yt.title}: views={val(yt.views)}, likes={val(yt.likes)}, comments={val(yt.comments)}")

    lines.append("")
    lines.append("### Social Media")
    for s in metrics.social:
        metric = s.reactions or s.likes or s.views
        status = "✓" if metric is not None else "⚠"
        lines.append(f"- {status} {s.platform} ({s.account}): {s.status}")

    lines.append("")
    lines.append("### Podcast")
    for p in metrics.podcast:
        status = "✓" if (p.plays or p.views or p.title) else "⚠"
        lines.append(f"- {status} {p.platform}: {p.title or 'title unavailable'}")

    output_path.write_text("\n".join(lines) + "\n")


def export_pdf(pptx_path: Path, pdf_path: Path) -> bool:
    try:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(pdf_path.parent),
                str(pptx_path),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
        generated = pptx_path.with_suffix(".pdf")
        if generated.exists() and generated != pdf_path:
            generated.rename(pdf_path)
        return pdf_path.exists()
    except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired) as exc:
        print(f"PDF export via LibreOffice failed: {exc}", file=sys.stderr)
        return False


def main() -> int:
    parser = argparse.ArgumentParser(description="Build Crownsmen Campaign Performance Report")
    parser.add_argument("config", nargs="?", default="links.yaml", help="Path to links.yaml")
    parser.add_argument("-o", "--output", default="output", help="Output directory")
    parser.add_argument("--metrics-file", help="Use pre-fetched metrics JSON instead of fetching")
    args = parser.parse_args()

    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.metrics_file:
        data = json.loads(Path(args.metrics_file).read_text())
        # Reconstruct CampaignMetrics from dict (simplified)
        from fetch_metrics import MetricResult
        metrics = CampaignMetrics(
            fetched_at=data["fetched_at"],
            campaign=data["campaign"],
            youtube=[MetricResult(**y) for y in data.get("youtube", [])],
            social=[MetricResult(**s) for s in data.get("social", [])],
            podcast=[MetricResult(**p) for p in data.get("podcast", [])],
            email=data.get("email", {}),
            google_ads=data.get("google_ads", {}),
            analytics=data.get("analytics", {}),
            missing=data.get("missing", []),
        )
    else:
        print("Fetching metrics from all URLs...")
        metrics = fetch_all(args.config)

    metrics_path = output_dir / "metrics.json"
    metrics_path.write_text(json.dumps(metrics_to_dict(metrics), indent=2))
    print(f"Metrics saved to {metrics_path}")

    campaign = metrics.campaign
    slug = f"Crownsmen_Partners_Campaign_Performance_Report__{campaign.get('episode', 'CS').replace(' ', '_')}_-_{campaign.get('client', 'Client').replace(' ', '_')}"
    pptx_path = output_dir / f"{slug}.pptx"
    pdf_path = output_dir / f"{slug}.pdf"

    print("Building PowerPoint presentation...")
    prs = build_presentation(metrics)
    prs.save(str(pptx_path))
    print(f"Presentation saved to {pptx_path}")

    missing_path = output_dir / "missing_data.md"
    write_missing_data(metrics, missing_path)
    print(f"Missing data report saved to {missing_path}")

    print("Exporting PDF...")
    if export_pdf(pptx_path, pdf_path):
        print(f"PDF saved to {pdf_path}")
    else:
        print("PDF export failed — PPTX is available", file=sys.stderr)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
